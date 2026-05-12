"""FastAPI web server for the RFP Summarizer UI."""

from __future__ import annotations

import asyncio
import json
import os
import re
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel

load_dotenv()

app = FastAPI(title="RFP Summarizer")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

STATIC_DIR = Path(__file__).parent / "static"


@app.get("/", response_class=HTMLResponse)
async def root():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/api/status")
async def status():
    return {
        "openai_key": bool(os.environ.get("OPENAI_API_KEY")),
        "sam_key": bool(os.environ.get("SAM_GOV_API_KEY")),
    }


# --------------------------------------------------------------------------- #
# Request models
# --------------------------------------------------------------------------- #


class SAMSearchRequest(BaseModel):
    query: str
    limit: int = 5
    posted_from: str | None = None
    posted_to: str | None = None
    model: str = "gpt-4o"


class SAMNoticeRequest(BaseModel):
    notice_id: str
    use_attachments: bool = False
    model: str = "gpt-4o"


class DownloadRequest(BaseModel):
    summary: dict[str, Any]


# --------------------------------------------------------------------------- #
# SSE helper
# --------------------------------------------------------------------------- #


def _sse(payload: dict) -> str:
    return f"data: {json.dumps(payload)}\n\n"


# --------------------------------------------------------------------------- #
# Endpoints
# --------------------------------------------------------------------------- #


@app.post("/api/search")
async def search_sam(req: SAMSearchRequest):
    """Stream SAM.gov search results via SSE — one summary per event."""
    if not os.environ.get("OPENAI_API_KEY"):
        raise HTTPException(400, "OPENAI_API_KEY not configured in .env")
    if not os.environ.get("SAM_GOV_API_KEY"):
        raise HTTPException(400, "SAM_GOV_API_KEY not configured in .env")

    def generate():
        try:
            from sam_gov import search_opportunities, opportunity_to_text
            from rfp_summarize import _summarize_text

            notices = search_opportunities(
                req.query,
                limit=req.limit,
                posted_from=req.posted_from,
                posted_to=req.posted_to,
            )

            yield _sse({"type": "total", "count": len(notices)})

            if not notices:
                yield _sse({"type": "done"})
                return

            for i, notice in enumerate(notices):
                title = notice.get("title", f"Opportunity {i + 1}")
                yield _sse({"type": "progress", "index": i, "title": title})

                text = opportunity_to_text(notice)
                summary = _summarize_text(
                    text, model=req.model, stream=False, output_format="none"
                )
                summary["_notice_id"] = notice.get("noticeId", "")
                summary["_sam_link"] = notice.get("uiLink", "")
                summary["_posted_date"] = notice.get("postedDate", "")
                summary["_source"] = "sam"

                yield _sse({"type": "summary", "index": i, "data": summary})

            yield _sse({"type": "done"})

        except Exception as exc:
            yield _sse({"type": "error", "message": str(exc)})

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/notice")
async def get_notice(req: SAMNoticeRequest):
    if not os.environ.get("OPENAI_API_KEY"):
        raise HTTPException(400, "OPENAI_API_KEY not configured in .env")

    def run():
        from rfp_summarize import summarize_from_sam_id
        return summarize_from_sam_id(
            req.notice_id,
            model=req.model,
            stream=False,
            output_format="none",
            use_attachments=req.use_attachments,
        )

    loop = asyncio.get_event_loop()
    summary = await loop.run_in_executor(None, run)
    summary["_notice_id"] = req.notice_id
    summary["_source"] = "sam"
    return {"summary": summary}


@app.post("/api/upload")
async def upload_file(
    file: UploadFile = File(...),
    model: str = Form(default="gpt-4o"),
):
    if not os.environ.get("OPENAI_API_KEY"):
        raise HTTPException(400, "OPENAI_API_KEY not configured in .env")

    suffix = Path(file.filename or "rfp.txt").suffix.lower()
    if suffix not in {".pdf", ".docx", ".doc", ".txt"}:
        raise HTTPException(400, f"Unsupported file type: {suffix}")

    content = await file.read()

    def run():
        from rfp_summarize import summarize_rfp
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            return summarize_rfp(
                tmp_path, model=model, stream=False, output_format="none"
            )
        finally:
            os.unlink(tmp_path)

    loop = asyncio.get_event_loop()
    summary = await loop.run_in_executor(None, run)
    summary["_source"] = "upload"
    summary["_filename"] = file.filename
    return {"summary": summary}


@app.post("/api/download/docx")
async def download_docx(req: DownloadRequest):
    """Generate and return a formatted .docx summary document."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise HTTPException(500, "python-docx is required: pip install python-docx")

    s = req.summary

    def _add_heading(doc, text, level=1):
        p = doc.add_heading(text, level=level)
        p.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
        return p

    def _add_kv(doc, label, value):
        if not value:
            return
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(10)
        p.add_run(str(value)).font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)

    def _add_bullets(doc, items):
        for item in (items or []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(str(item)).font.size = Pt(10)
            p.paragraph_format.space_after = Pt(2)

    def _shade_row(row, hex_color):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hex_color)
            tcPr.append(shd)

    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title
    title_p = doc.add_heading(s.get("title") or "RFP Summary", 0)
    title_p.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
    title_p.runs[0].font.size = Pt(18)

    if s.get("issuing_organization"):
        sub = doc.add_paragraph(s["issuing_organization"])
        sub.runs[0].font.size = Pt(11)
        sub.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
        sub.paragraph_format.space_after = Pt(4)

    doc.add_paragraph()

    # Key info table
    _add_heading(doc, "Key Information", 2)
    info_rows = [
        ("Solicitation Number", s.get("rfp_number")),
        ("Issue Date", s.get("issue_date")),
        ("Proposal Due Date", s.get("due_date")),
        ("Questions Deadline", s.get("questions_deadline")),
        ("Estimated Budget", s.get("budget", {}).get("amount") if isinstance(s.get("budget"), dict) else s.get("budget")),
        ("Budget Notes", s.get("budget", {}).get("notes") if isinstance(s.get("budget"), dict) else None),
        ("Contract Type", s.get("contract_type")),
        ("Contract Period", s.get("contract_period")),
        ("Submission Format", s.get("submission_format")),
    ]
    info_rows = [(k, v) for k, v in info_rows if v]
    if info_rows:
        tbl = doc.add_table(rows=len(info_rows), cols=2)
        tbl.style = "Table Grid"
        for i, (k, v) in enumerate(info_rows):
            tbl.rows[i].cells[0].text = k
            tbl.rows[i].cells[1].text = str(v)
            tbl.rows[i].cells[0].paragraphs[0].runs[0].bold = True
            tbl.rows[i].cells[0].paragraphs[0].runs[0].font.size = Pt(9)
            tbl.rows[i].cells[1].paragraphs[0].runs[0].font.size = Pt(9)
            _shade_row(tbl.rows[i], "F8FAFC" if i % 2 == 0 else "FFFFFF")
        doc.add_paragraph()

    # Overview
    if s.get("project_overview"):
        _add_heading(doc, "Project Overview", 2)
        p = doc.add_paragraph(s["project_overview"])
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    # Scope of work
    if s.get("scope_of_work"):
        _add_heading(doc, "Scope of Work", 2)
        _add_bullets(doc, s["scope_of_work"])
        doc.add_paragraph()

    # Requirements
    mandatory = (s.get("requirements") or {}).get("mandatory", [])
    preferred = (s.get("requirements") or {}).get("preferred", [])
    if mandatory or preferred:
        _add_heading(doc, "Requirements", 2)
        if mandatory:
            p = doc.add_paragraph("Mandatory Requirements")
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
            _add_bullets(doc, mandatory)
        if preferred:
            p = doc.add_paragraph("Preferred Qualifications")
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.color.rgb = RGBColor(0x1d, 0x4e, 0xd8)
            _add_bullets(doc, preferred)
        doc.add_paragraph()

    # Evaluation criteria
    eval_criteria = s.get("evaluation_criteria", [])
    if eval_criteria:
        _add_heading(doc, "Evaluation Criteria", 2)
        tbl = doc.add_table(rows=len(eval_criteria) + 1, cols=2)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0]
        hdr.cells[0].text = "Criterion"
        hdr.cells[1].text = "Weight"
        _shade_row(hdr, "1E3A5F")
        for cell in hdr.cells:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell.paragraphs[0].runs[0].font.size = Pt(9)
        for i, c in enumerate(eval_criteria):
            row = tbl.rows[i + 1]
            row.cells[0].text = c.get("criterion", "")
            row.cells[1].text = c.get("weight", "")
            for cell in row.cells:
                cell.paragraphs[0].runs[0].font.size = Pt(9)
            _shade_row(row, "F8FAFC" if i % 2 == 0 else "FFFFFF")
        doc.add_paragraph()

    # Eligibility
    if s.get("eligibility"):
        _add_heading(doc, "Eligibility / Set-Aside", 2)
        _add_bullets(doc, s["eligibility"])
        doc.add_paragraph()

    # Key dates
    key_dates = s.get("key_dates", [])
    if key_dates:
        _add_heading(doc, "Key Dates", 2)
        tbl = doc.add_table(rows=len(key_dates) + 1, cols=2)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0]
        hdr.cells[0].text = "Event"
        hdr.cells[1].text = "Date"
        _shade_row(hdr, "1E3A5F")
        for cell in hdr.cells:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell.paragraphs[0].runs[0].font.size = Pt(9)
        for i, d in enumerate(key_dates):
            row = tbl.rows[i + 1]
            row.cells[0].text = d.get("event", "")
            row.cells[1].text = d.get("date", "")
            for cell in row.cells:
                cell.paragraphs[0].runs[0].font.size = Pt(9)
            _shade_row(row, "F8FAFC" if i % 2 == 0 else "FFFFFF")
        doc.add_paragraph()

    # Submission requirements
    if s.get("submission_requirements"):
        _add_heading(doc, "Submission Requirements", 2)
        _add_bullets(doc, s["submission_requirements"])
        doc.add_paragraph()

    # Contact
    contact = s.get("contact") or {}
    if any(contact.values()):
        _add_heading(doc, "Point of Contact", 2)
        _add_kv(doc, "Name", contact.get("name"))
        _add_kv(doc, "Email", contact.get("email"))
        _add_kv(doc, "Phone", contact.get("phone"))
        doc.add_paragraph()

    # Attachments
    if s.get("attachments"):
        _add_heading(doc, "Attachments", 2)
        _add_bullets(doc, s["attachments"])
        doc.add_paragraph()

    # Red flags
    red_flags = s.get("red_flags", [])
    if red_flags:
        _add_heading(doc, "⚠ Red Flags", 2)
        for flag in red_flags:
            p = doc.add_paragraph()
            run = p.add_run(f"⚠  {flag}")
            run.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
            run.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(4)

    # SAM.gov link
    sam_link = s.get("_sam_link") or s.get("sam_gov_link")
    if sam_link:
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run(f"SAM.gov: {sam_link}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xeb)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    raw = s.get("rfp_number") or s.get("title") or "summary"
    safe = re.sub(r"[^a-z0-9]", "_", raw.lower())[:40]

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="rfp_summary_{safe}.docx"'},
    )

